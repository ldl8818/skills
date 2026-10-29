"""Focused regressions for chart fidelity and new asset/template entry points."""
import argparse
import copy
import io
import contextlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
import subprocess

from PIL import Image, ImageChops
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import brand_kit
import chartcheck
import photo_assets
import template_kit
import build_deck
import common
import layouts
import extract_deck
import outline_check
import render_lite
from unittest import mock


class ChartPreviewTests(unittest.TestCase):
    def preview(self, spec):
        renderer = render_lite.Painter(common.Style({}), Path("."), 40)
        image = Image.new("RGBA", (renderer.W, renderer.H), "white")
        rects, lines = [], []
        original_rect, original_line = renderer.rect, renderer.line
        def rect(im, element, oval=False):
            rects.append((dict(element), oval))
            return original_rect(im, element, oval)
        def line(im, element):
            lines.append(dict(element))
            return original_line(im, element)
        with mock.patch.object(renderer, "rect", side_effect=rect), mock.patch.object(renderer, "line", side_effect=line):
            renderer.chart(image, {"chart": spec, "x": 1, "y": 1, "w": 8, "h": 4})
        return image, rects, lines

    def test_scatter_uses_each_series_xy_points_without_categories(self):
        spec = {"type": "scatter", "series": [{"name": "A", "x": [1, 3], "values": [-2, 4]},
                                                {"name": "B", "x": [2], "values": [1]}]}
        image, rects, lines = self.preview(spec)
        dots = [e for e, oval in rects if oval]
        self.assertEqual(len(dots), 3)
        self.assertLess(dots[0]["x"], dots[2]["x"])
        self.assertLess(dots[2]["x"], dots[1]["x"])
        self.assertGreater(dots[0]["y"], dots[2]["y"])
        self.assertGreater(dots[2]["y"], dots[1]["y"])
        self.assertTrue(lines)
        self.assertIsNotNone(ImageChops.difference(image.convert("RGB"), Image.new("RGB", image.size, "white")).getbbox())
        for dot in dots:
            center = (round((dot["x"] + dot["w"] / 2) * 80), round((dot["y"] + dot["h"] / 2) * 80))
            self.assertNotEqual(image.getpixel(center)[:3], (255, 255, 255))

    def test_mixed_stacks_share_zero_and_accumulate_each_sign(self):
        spec = {"type": "stacked", "categories": ["Q1"], "series": [
            {"name": "A", "values": [-10]}, {"name": "B", "values": [6]},
            {"name": "C", "values": [-5]}, {"name": "D", "values": [4]}]}
        image, rects, lines = self.preview(spec)
        bars = [e for e, oval in rects if e["h"] > 0.2 and not oval]
        self.assertEqual(len(bars), 4)
        zero = next(e["y1"] for e in lines if e["y1"] == e["y2"])
        negative, positive, negative_next, positive_next = bars
        self.assertAlmostEqual(negative["y"], zero)
        self.assertAlmostEqual(positive["y"] + positive["h"], zero)
        self.assertAlmostEqual(negative_next["y"], negative["y"] + negative["h"])
        self.assertAlmostEqual(positive_next["y"] + positive_next["h"], positive["y"])
        self.assertAlmostEqual(negative["h"] / positive["h"], 10 / 6)

    def test_negative_column_bar_and_line_stay_in_the_plot(self):
        for kind in ("column", "bar", "line"):
            with self.subTest(kind=kind):
                _, rects, lines = self.preview({"type": kind, "categories": ["A", "B"],
                    "series": [{"name": "Revenue", "values": [-10, 5]}]})
                self.assertTrue(rects)
                for e, oval in rects:
                    self.assertGreater(e["w"], 0)
                    self.assertGreater(e["h"], 0)
                    self.assertGreaterEqual(e["y"], 1)
                    self.assertLessEqual(e["y"] + e["h"], 5)

    def test_horizontal_bars_extend_to_opposite_sides_of_zero(self):
        _, rects, lines = self.preview({"type": "bar", "categories": ["A", "B"],
            "series": [{"name": "Revenue", "values": [-10, 5]}]})
        bars = [e for e, oval in rects if not oval]
        zero = next(e["x1"] for e in lines if e["x1"] == e["x2"])
        self.assertAlmostEqual(bars[0]["x"] + bars[0]["w"], zero)
        self.assertAlmostEqual(bars[1]["x"], zero)
        self.assertAlmostEqual(bars[0]["w"] / bars[1]["w"], 2)

    def test_invalid_chart_prevents_a_successful_preview(self):
        painter = render_lite.Painter(common.Style({}), Path("."), 40)
        image = Image.new("RGBA", (painter.W, painter.H), "white")
        for chart in ({"type": "scatter", "series": []}, {"type": "bubble", "categories": ["A"],
                      "series": [{"values": [3]}]}):
            with self.subTest(chart=chart), self.assertRaises(ValueError):
                painter.draw(image, {"n": 1, "elements": [{"t": "chart", "x": 1, "y": 1,
                                                          "w": 8, "h": 4, "chart": chart}]})



class AliasFidelityTests(unittest.TestCase):
    def qa(self, slide, layout=None, drop_text=None, ghost=False):
        with tempfile.TemporaryDirectory() as td:
            work = Path(td)
            slide = {"n": 1, **slide}
            deck = {"meta": {}, "slides": [slide]}
            (work / "deck.json").write_text(json.dumps(deck))
            (work / "style.json").write_text(json.dumps({"decor": {"closing_dark": False}}))
            layout = layout or layouts.TYPE_TO_LAYOUTS[slide["type"]][0]
            (work / "plan.json").write_text(json.dumps({"slides": [{"n": 1, "layout": layout}], "assets": []}))
            scripts = Path(__file__).resolve().parents[1] / "scripts"
            result = subprocess.run([sys.executable, str(scripts / "build_deck.py"), str(work)],
                                    text=True, capture_output=True)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            if drop_text is not None:
                path = work / "out/beautified.pptx"
                prs = Presentation(path)
                shape = next(s for s in prs.slides[0].shapes if s.has_text_frame and s.text == drop_text)
                shape._element.getparent().remove(shape._element)
                prs.save(path)
            result = subprocess.run([sys.executable, str(scripts / "qa_check.py"), str(work)],
                                    text=True, capture_output=True)
            report = json.loads((work / "out/qa_report.json").read_text())
            fidelity = [i for i in report["issues"] if i["check"] == "fidelity" and i["level"] == "ERROR"]
            return report["ghost_deck"] if ghost else fidelity

    def test_title_fallback_aliases_do_not_require_duplicate_shapes(self):
        for typ, field in (("closing", "headline"), ("question", "question"), ("big_quote", "quote")):
            for fallback in ("Same message", "Inactive title"):
                slide = {"type": typ, "title": fallback, field: "Same message", "body": "Visible body"}
                if typ == "big_quote":
                    slide.pop("body")
                with self.subTest(type=typ, fallback=fallback):
                    self.assertEqual(self.qa(slide), [])
                    self.assertEqual(self.qa(slide, ghost=True), ["1. Same message"])
        for typ in ("closing", "question"):
            self.assertEqual(self.qa({"type": typ, "title": "Fallback message"}), [])
        self.assertTrue(self.qa({"type": "big_quote", "title": "Unrendered title"}))

    def test_alternative_fields_follow_the_effective_layout(self):
        cases = [
            {"type": "hub", "title": "Hub", "center": {"title": "Center"}, "center_title": "Inactive center",
             "items": [{"title": "One", "text": "Alpha"}, {"title": "Two", "text": "Beta"}]},
            {"type": "stats", "title": "Stats", "body": "Visible body",
             "detail": {"title": "Inactive detail", "body": "Inactive explanation"},
             "items": [{"value": "13", "label": "Orders"}]},
            {"type": "explain", "title": "Explain", "blocks": [{"text": "Inactive blocks"}],
             "columns": [{"blocks": [{"text": "Visible column"}]}]},
        ]
        for slide in cases:
            with self.subTest(type=slide["type"]):
                self.assertEqual(self.qa(slide), [])
        # A manual layout override renders both fields as real content.
        slide = {"type": "closing", "title": "Same message", "headline": "Same message"}
        self.assertEqual(self.qa(slide, "statement_hero"), [])
        self.assertTrue(self.qa(slide, "statement_hero", drop_text="Same message"))

    def test_real_repeated_body_and_title_remain_counted(self):
        slide = {"type": "statement", "title": "Same message", "headline": "Same message"}
        self.assertEqual(self.qa(slide), [])
        self.assertTrue(self.qa(slide, drop_text="Same message"))
        repeated = {"type": "list", "title": "Repeat body", "items": [{"text": "Alpha"}, {"text": "Alpha"}]}
        self.assertEqual(self.qa(repeated), [])
        self.assertTrue(self.qa(repeated, drop_text="Alpha"))


class ReviewTests(unittest.TestCase):
    def test_custom_chart_labels_refused_without_replacing_work(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "source.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            data = CategoryChartData()
            data.categories = ["A", "B"]
            data.add_series("Revenue", [10, 20])
            chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                    Inches(1), Inches(1), Inches(6), Inches(4), data).chart
            chart.plots[0].has_data_labels = True
            chart.series[0].points[0].data_label.text_frame.text = "Custom label"
            prs.save(source)
            work = root / "work"
            work.mkdir()
            (work / "source_text.json").write_text("KEEP")
            from unittest import mock
            with mock.patch.object(sys, "argv", ["extract_deck.py", str(source), str(work), "--no-render"]):
                with self.assertRaises(SystemExit):
                    extract_deck.main()
            self.assertEqual((work / "source_text.json").read_text(), "KEEP")

    def test_render_ids_cover_new_shapes_and_allow_later_proxy_edits(self):
        with tempfile.TemporaryDirectory() as td:
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            old = slide.shapes.add_textbox(0, 0, Inches(1), Inches(1)).shape_id
            builder = build_deck.Builder(Path(td), common.Style({}), {}, {}, 13.333, 7.5)
            elements = [layouts.R(1, 1, 1, 1, fill="primary"),
                        layouts.T("Editable", 2, 2, 4, 1, 20),
                        {"t": "unknown"}]
            rendered = builder.render(slide, elements)
            ids = [i for _, group in rendered for i in group]
            self.assertEqual(set(ids), {s.shape_id for s in slide.shapes} - {old})
            self.assertEqual([len(group) for _, group in rendered], [1, 1, 0])
            added = prs.slides[0].shapes.add_textbox(0, 0, Inches(1), Inches(1)).shape_id
            self.assertNotIn(added, ids + [old])
            out = Path(td) / "render.pptx"
            prs.save(out)
            self.assertEqual(len(Presentation(out).slides[0].shapes), 4)

    def test_manual_refresh_removes_stale_pages_after_success_only(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            src = root / "page.png"
            Image.new("RGB", (50, 50), "white").save(src)
            work = root / "work"
            args = dict(work=work, dpi=100)
            brand_kit.cmd_pages(argparse.Namespace(**args, src=[src, src]))
            self.assertEqual(len(list((work / "brand/pages").glob("p*.png"))), 2)
            with self.assertRaises(FileNotFoundError):
                brand_kit.cmd_pages(argparse.Namespace(**args, src=[root / "missing.png"]))
            self.assertEqual(len(list((work / "brand/pages").glob("p*.png"))), 2)
            brand_kit.cmd_pages(argparse.Namespace(**args, src=[src]))
            self.assertEqual(len(list((work / "brand/pages").glob("p*.png"))), 1)

    def test_chart_checks_each_repeated_category_and_series(self):
        for duplicate_series in (False, True):
            spec = {"type": "column", "categories": ["Q1", "Q1"], "series": [
                {"name": "Revenue", "values": [10, 20]}]}
            if duplicate_series:
                spec["series"].append({"name": "Revenue", "values": [30, 40]})
            ref = chartcheck.from_spec(spec)
            self.assertFalse(chartcheck.compare(ref, copy.deepcopy(ref))[0])
            wrong = copy.deepcopy(ref)
            wrong["series"][0]["values"][0] = 999
            self.assertTrue(chartcheck.compare(ref, wrong)[0], "every occurrence must be checked")

    def test_photo_drop_rejects_escape_before_manifest_mutation(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            work = root / "work"
            (work / "assets").mkdir(parents=True)
            manifest = work / "assets/manifest.json"
            manifest.write_text('{"assets": []}')
            outside = root / "outside.png"
            outside.write_bytes(b"KEEP")
            before = manifest.read_bytes()
            with self.assertRaises(ValueError):
                photo_assets.cmd_drop(argparse.Namespace(work=work, id="../../outside"), {})
            self.assertEqual(outside.read_bytes(), b"KEEP")
            self.assertEqual(manifest.read_bytes(), before)
            photo_assets.cmd_drop(argparse.Namespace(work=work, id="hero_01"), {})

    def test_brand_crop_rejects_escape_and_symlink(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            work = root / "work"
            (work / "brand").mkdir(parents=True)
            source = root / "source.png"
            Image.new("RGB", (20, 20), "blue").save(source)
            outside = root / "outside.png"
            outside.write_bytes(b"KEEP")
            (work / "brand/linked.png").symlink_to(outside)
            args = dict(work=work, page=source, x0=0, y0=0, x1=1, y1=1, trim_bottom=0, clear=False)
            for name in ("../../outside", "linked"):
                with self.subTest(name=name), self.assertRaises(ValueError):
                    brand_kit.cmd_crop(argparse.Namespace(**args, name=name))
                self.assertEqual(outside.read_bytes(), b"KEEP")
            brand_kit.cmd_crop(argparse.Namespace(**args, name="logo"))
            self.assertTrue((work / "brand/logo.png").is_file())

    def test_template_read_cannot_destroy_source_via_output_alias(self):
        with tempfile.TemporaryDirectory() as td:
            work = Path(td)
            (work / "template").mkdir()
            source = work / "template/template.pptx"
            prs = Presentation()
            prs.slides.add_slide(prs.slide_layouts[6])
            prs.save(source)
            before = source.read_bytes()
            with self.assertRaises(ValueError):
                template_kit.cmd_read(argparse.Namespace(work=work, template=source, no_render=True, base=None))
            self.assertEqual(source.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
