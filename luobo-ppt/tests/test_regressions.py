"""Offline regressions. Every generated file lives in a TemporaryDirectory."""
import contextlib
import importlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

from lxml import etree
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / "scripts"))
import common
import omml
import layouts
import extract_deck
import build_deck
import qa_check
import render_lite
import make_assets
import imagegen
import pipeline
import plan_deck
import style_previews
import render
import compare_sheet
import analyze_source


def call(module, *args):
    output = io.StringIO()
    with mock.patch.object(sys, "argv", [module.__file__, *map(str, args)]), contextlib.redirect_stdout(output):
        try:
            rc = module.main()
        except SystemExit as exc:
            rc = exc.code
    return rc or 0, output.getvalue()


def put(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


class Isolated(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="luobo-test-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.addCleanup(mock.patch.stopall)
        mock.patch.dict(os.environ, {"HOME": str(self.root / "home"), "PYTHONDONTWRITEBYTECODE": "1"}).start()
        mock.patch.object(common, "ICON_CACHE", self.root / "icons").start()
        mock.patch.object(common, "SETUP_CACHE", self.root / "setup.json").start()
        mock.patch.object(make_assets, "find_config", return_value={"provider": "none"}).start()
        mock.patch.object(imagegen, "find_config", return_value={"provider": "none"}).start()
        mock.patch("socket.socket.connect", side_effect=AssertionError("network forbidden")).start()
        self.cwd = Path.cwd()
        os.chdir(self.root)
        self.addCleanup(os.chdir, self.cwd)
        self.style = common.Style({"palette": {"bg_bottom": "FFFFFF", "text": "000000"}})

    def build(self, text, source=None, elements=None, work=None):
        work = work or self.root / "work"
        put(work / "style.json", self.style.data)
        put(work / "deck.json", {"meta": {}, "slides": [{"n": 1, "type": "list", "title": text}]})
        put(work / "plan.json", {"slides": [{"n": 1, "layout": "free", "elements":
            elements if elements is not None else [layouts.T(text, 1, 2, 10, 1, 24)]}], "assets": []})
        if source is not None:
            put(work / "source_text.json", {"slides": [{"n": 1, "units":
                [{"id": str(i), "text": s} for i, s in enumerate(source)]}]})
        self.assertEqual(call(build_deck, work)[0], 0)
        return work

    def qa(self, work):
        rc, _ = call(qa_check, work)
        return rc, json.loads((work / "out/qa_report.json").read_text())

    def preview(self, work):
        self.assertEqual(call(render_lite, work, "--dpi", 40)[0], 0)
        with Image.open(work / "out/png/slide-01.png") as im:
            return im.copy()


class FidelityTests(Isolated):
    def test_math_extract_build_qa_and_pixels(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
        cases = [("x", "2", "x\u00b2"), ("x", "n", "x\u207f"),
                 ("x", "y", "x^(y)"), ("a+b", "2", "(a+b)\u00b2")]
        for i, (base, exponent, _) in enumerate(cases):
            para = box.text_frame.paragraphs[0] if not i else box.text_frame.add_paragraph()
            xml = (f'<a14:m xmlns:a14="{omml.NS["a14"]}" xmlns:m="{omml.NS["m"]}">'
                   f'<m:oMath><m:sSup><m:e><m:r><m:t>{base}</m:t></m:r></m:e>'
                   f'<m:sup><m:r><m:t>{exponent}</m:t></m:r></m:sup></m:sSup></m:oMath></a14:m>')
            para._p.append(etree.fromstring(xml.encode()))
        original = self.root / "original.pptx"
        prs.save(original)
        original_bytes = original.read_bytes()
        work = self.root / "extracted"
        self.assertEqual(call(extract_deck, original, work, "--no-render")[0], 0)
        units = json.loads((work / "source_text.json").read_text())["slides"][0]["units"]
        self.assertEqual([u["text"] for u in units], [c[2] for c in cases])
        self.assertEqual(original.read_bytes(), original_bytes)
        text = "\n".join(c[2] for c in cases)
        self.build(text, elements=[layouts.T(text, 1, 2, 10, 3, 24)], work=work)
        self.assertEqual(self.qa(work)[0], 0)
        preview = self.preview(work)
        crop = preview.crop((40, 80, 400, 180))
        self.assertIsNotNone(ImageChops.difference(crop, Image.new("RGB", crop.size, "white")).getbbox())

    def test_unknown_subscript_and_combined(self):
        for tag, tail, expected in [
            ("sSub", "<m:sub><m:r><m:t>q</m:t></m:r></m:sub>", "x_(q)"),
            ("sSubSup", "<m:sub><m:r><m:t>q</m:t></m:r></m:sub>"
             "<m:sup><m:r><m:t>y</m:t></m:r></m:sup>", "x_(q)^(y)"),
        ]:
            xml = f'<m:{tag} xmlns:m="{omml.NS["m"]}"><m:e><m:r><m:t>x</m:t></m:r></m:e>{tail}</m:{tag}>'
            self.assertEqual(omml.lin(etree.fromstring(xml.encode())), expected)

    def test_negative_and_multiplicity_fail_but_split_text_passes(self):
        cases = [
            (["-1"], "1", True), (["-1"], "-1", False),
            (["repeat", "repeat"], "repeat", True),
            (["repeat", "repeat"], "repeat\nrepeat", False),
            (["alpha beta"], "alpha\nbeta", False),
            (["- bullet"], "\u2022 bullet", False),
            (["a", "abc"], "abc", True),
            (["a", "abc"], "a\nabc", False),
        ]
        for i, (source, text, failed) in enumerate(cases):
            with self.subTest(source=source, text=text):
                work = self.build(text, source, work=self.root / str(i))
                self.assertEqual(bool(self.qa(work)[0]), failed)
                preview = self.preview(work)
                crop = preview.crop((40, 80, 220, 120))
                self.assertIsNotNone(ImageChops.difference(crop, Image.new("RGB", crop.size, "white")).getbbox())

    def test_chart_refused_before_mutating_work_or_original(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ["Alpha", "Beta"]
        data.add_series("Revenue", [123, 456])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(8), Inches(4), data)
        src = self.root / "chart.pptx"
        prs.save(src)
        original = src.read_bytes()
        work = self.root / "existing"
        work.mkdir()
        (work / "source_text.json").write_text("KEEP")
        rc, output = call(extract_deck, src, work, "--no-render")
        self.assertTrue(rc)
        self.assertIn("chart", (str(rc) + output).lower())
        self.assertIn("1", str(rc) + output)
        self.assertEqual(src.read_bytes(), original)
        self.assertEqual((work / "source_text.json").read_text(), "KEEP")
        self.assertFalse((work / "source.pptx").exists())

    def test_text_and_table_extract_normally(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Text"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
        table.cell(0, 0).text = "Alpha"
        table.cell(0, 1).text = "123"
        src = self.root / "normal.pptx"
        prs.save(src)
        work = self.root / "normal"
        self.assertEqual(call(extract_deck, src, work, "--no-render")[0], 0)
        units = json.loads((work / "source_text.json").read_text())["slides"][0]["units"]
        self.assertEqual([u["text"] for u in units], ["Text", "Alpha", "123"])

    def test_long_token_fit_overflow_and_preview(self):
        text = "W" * 120
        size = common.fit_size(text, 1, 1, 36, 12, max_lines=1)
        self.assertLess(size, 36)
        work = self.build(text, elements=[layouts.T(text, 1, 2, 1, 1, 36)])
        rc, report = self.qa(work)
        self.assertNotEqual(rc, 0)
        self.assertTrue(any(i["check"] == "overflow" for i in report["issues"]))
        text = "W" * 24
        e = layouts.T(text, 1, 2, 1.2, 3, 20)
        work = self.build(text, elements=[e], work=self.root / "fitting")
        self.assertEqual(self.qa(work)[0], 0)
        preview = self.preview(work)
        inside = preview.crop((40, 80, 88, 200))
        outside = preview.crop((92, 80, 160, 200))
        self.assertIsNotNone(ImageChops.difference(inside, Image.new("RGB", inside.size, "white")).getbbox())
        self.assertIsNone(ImageChops.difference(outside, Image.new("RGB", outside.size, "white")).getbbox())
        self.assertEqual("".join(common.wrap_lines(text, 20, 80)), text)
        self.assertEqual(common.wrap_lines("short text", 12, 300), ["short text"])


class AssetTests(Isolated):
    def test_size_rejected_before_provider_or_retry(self):
        gen = imagegen.ImageGen({"provider": "command", "command": "never"})
        for bad in [True, False, 0, -1, 1.5, "100", "1; echo injected #", None]:
            for size in [(bad, 10), (10, bad)]:
                with self.subTest(size=size), mock.patch.object(gen, "_gen_command") as provider, \
                     mock.patch.object(imagegen.time, "sleep") as sleep:
                    with self.assertRaises(ValueError):
                        gen.generate("test", *size, out=str(self.root / "out.png"))
                    provider.assert_not_called()
                    sleep.assert_not_called()

    def test_command_prompt_and_output_are_single_shell_arguments(self):
        runner = self.root / "record.py"
        runner.write_text("import json,sys\nfrom pathlib import Path\n"
                          "Path(sys.argv[1]).write_text(json.dumps(sys.argv[2:]))\n", encoding="utf-8")
        import shlex
        record = self.root / "argv.json"
        command = f"{shlex.quote(sys.executable)} {shlex.quote(str(runner))} {shlex.quote(str(record))} {{w}} {{h}} {{out}} {{prompt}} {{prompt_file}}"
        gen = imagegen.ImageGen({"provider": "command", "command": command, "retries": 0})
        out = self.root / "image $(touch INJECTED); ' quote.png"
        prompt = "prompt $(touch INJECTED); ' \" `touch INJECTED`"
        Image.new("RGB", (2, 2), "red").save(out)
        image, _ = gen.generate(prompt, 2, 2, out=str(out))
        image.close()
        self.assertEqual(json.loads(record.read_text()), ["2", "2", str(out), prompt, str(out.with_suffix(".prompt.txt"))])
        self.assertFalse((self.root / "INJECTED").exists())

    def test_asset_ids_and_symlinks_cannot_escape(self):
        work = self.root / "work"
        put(work / "style.json", {})
        for child in ["assets", "prompts"]:
            (work / child).mkdir()
        outside = self.root / "outside.md"
        outside.write_text("KEEP")
        for aid in ["../../outside", str(self.root / "absolute"), r"..\outside", "..", "", "bad/name"]:
            with self.subTest(aid=aid):
                asset = {"id": aid, "kind": "hero", "prompt": "new"}
                with self.assertRaises(ValueError):
                    make_assets.prompt_file(work, asset, self.style, "hero", rewrite=True)
        (work / "prompts/linked.md").symlink_to(outside)
        with self.assertRaises(ValueError):
            make_assets.prompt_file(work, {"id": "linked", "prompt": "new"}, self.style, "hero", rewrite=True)
        self.assertEqual(outside.read_text(), "KEEP")
        _, path = make_assets.prompt_file(work, {"id": "hero_01", "prompt": "normal"}, self.style, "hero")
        self.assertTrue(path.is_file())
        for assets in [[{"id": "../../outside", "kind": "hero"}],
                       [{"id": "sheet_01", "kind": "sheet", "items": [{"id": "../../outside"}]}]]:
            put(work / "plan.json", {"assets": assets})
            with mock.patch.object(make_assets, "procedural") as generate:
                with self.assertRaises(ValueError):
                    call(make_assets, work, "--procedural")
                generate.assert_not_called()

    def test_manifest_relative_work_and_absolute_work_preview_pixels(self):
        for relative in [True, False]:
            work = Path("relative") if relative else self.root / "absolute"
            (work / "assets").mkdir(parents=True)
            Image.new("RGB", (8, 8), "red").save(work / "assets/bg.jpg")
            put(work / "assets/manifest.json", {"assets": [{"id": "bg_content", "path": str(work / "assets/bg.jpg")}]})
            self.build("", elements=[], work=work)
            prs = Presentation(work / "out/beautified.pptx")
            self.assertEqual(sum(s.shape_type == 13 for s in prs.slides[0].shapes), 1)
            pixel = self.preview(work).getpixel((5, 5))
            self.assertGreater(pixel[0], 240)
            self.assertLess(pixel[1], 10)


class WorkflowTests(Isolated):
    def test_pipeline_real_cli_status_and_preview(self):
        work = self.build("Alpha", source=["Alpha"])
        command = [sys.executable, "-B", str(SKILL / "scripts/pipeline.py"),
                   str(work), "--lite", "--procedural", "--dpi", "40"]
        good = subprocess.run(command, capture_output=True, text=True, timeout=30)
        self.assertEqual(good.returncode, 0, good.stdout + good.stderr)
        self.assertTrue((work / "out/sheet.png").is_file())
        put(work / "source_text.json", {"slides": [{"n": 1, "units": [{"id": "1", "text": "Missing"}]}]})
        bad = subprocess.run(command, capture_output=True, text=True, timeout=30)
        self.assertNotEqual(bad.returncode, 0)
        self.assertIn("QA has ERRORS", bad.stdout)

    def test_pipeline_failure_status(self):
        for failed in [None, "qa_check.py", "review.py", "render_lite.py"]:
            def run(args):
                return 7 if Path(args[0]).name == failed else 0
            with self.subTest(failed=failed), mock.patch.object(pipeline, "run", side_effect=run):
                rc, _ = call(pipeline, self.root, "--lite")
                self.assertEqual(bool(rc), failed is not None)

    def test_previews_fallback_and_image_rebasing(self):
        work = self.root / "preview"
        (work / "media").mkdir(parents=True)
        Image.new("RGB", (12, 12), "blue").save(work / "media/shot.png")
        deck = {"meta": {}, "slides": [{"n": 1, "type": "cover", "title": "Cover"},
                {"n": 2, "type": "screenshot", "title": "Screen", "image": "media/shot.png",
                 "stamps": [{"image": "media/shot.png", "x": 1, "y": 1, "w": 1, "h": 1}]}]}
        seen = []
        def run(args):
            seen.append(Path(args[0]).name)
            name = Path(args[0]).stem
            if name == "render":
                return subprocess.CompletedProcess(args, 1, "", "missing soffice")
            if name == "make_assets":
                # No illustration generation is needed to exercise the actual build/preview path.
                return subprocess.CompletedProcess(args, 0, "", "")
            rc, output = call(importlib.import_module(name), *args[1:])
            return subprocess.CompletedProcess(args, rc, output, "")
        with mock.patch.object(style_previews, "run", side_effect=run):
            files = style_previews.build_one(work, "mint-fresh", deck)
        self.assertIn("render_lite.py", seen)
        self.assertEqual(len(files), 2)
        prs = Presentation(work / "previews/mint-fresh/out/beautified.pptx")
        self.assertGreaterEqual(sum(s.shape_type == 13 for s in prs.slides[1].shapes), 2)
        with Image.open(files[1]) as im:
            self.assertTrue(any(b > 200 and r < 30 and g < 30 for _, (r, g, b) in
                                im.convert("RGB").getcolors(im.width * im.height)))
        self.assertEqual(deck["slides"][1]["image"], "media/shot.png")

    def test_preview_normal_renderer_still_used(self):
        seen = []
        def run(args):
            seen.append(Path(args[0]).name)
            if Path(args[0]).name == "render.py":
                target = Path(args[2])
                target.mkdir()
                Image.new("RGB", (8, 8), "white").save(target / "slide-01.png")
            return subprocess.CompletedProcess(args, 0, "", "")
        with mock.patch.object(style_previews, "run", side_effect=run):
            files = style_previews.build_one(self.root, "mint-fresh", style_previews.SAMPLE)
        self.assertEqual(len(files), 1)
        self.assertNotIn("render_lite.py", seen)

    def test_level_changes_preserve_only_compatible_manual_edits(self):
        work = self.root / "plan"
        put(work / "style.json", {})
        put(work / "deck.json", {"meta": {}, "slides": [{"n": 1, "type": "cover", "title": "Cover"}]})
        call(plan_deck, work, "--level", "full")
        plan = json.loads((work / "plan.json").read_text())
        hero = next(a for a in plan["assets"] if a["kind"] == "hero")
        hero["prompt"] = "hand edited prompt"
        plan["slides"][0]["relation"] = "hand edited relation"
        put(work / "plan.json", plan)
        call(plan_deck, work, "--level", "full")
        plan = json.loads((work / "plan.json").read_text())
        self.assertEqual(next(a for a in plan["assets"] if a["kind"] == "hero")["prompt"], "hand edited prompt")
        call(plan_deck, work, "--level", "light")
        plan = json.loads((work / "plan.json").read_text())
        self.assertFalse(plan["slides"][0].get("hero"))
        self.assertFalse(any(a["kind"] in ("hero", "sheet") for a in plan["assets"]))
        self.assertEqual(plan["slides"][0]["relation"], "hand edited relation")
        call(plan_deck, work, "--level", "full")
        self.assertTrue(json.loads((work / "plan.json").read_text())["slides"][0].get("hero"))

    def test_render_cleanup_all_paths(self):
        scratch = self.root / "scratch"
        scratch.mkdir()
        src = self.root / "source.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(src)
        with mock.patch.object(tempfile, "tempdir", str(scratch)):
            for mode in ["missing", "timeout", "failed", "success"]:
                def convert(cmd, **kwargs):
                    if mode == "timeout":
                        raise subprocess.TimeoutExpired(cmd, 1)
                    if mode == "success":
                        Path(cmd[-1]).with_suffix(".pdf").write_bytes(b"synthetic pdf")
                    return subprocess.CompletedProcess(cmd, 0)
                with self.subTest(mode=mode), mock.patch.object(render, "find_soffice", return_value=None if mode == "missing" else "fake"), mock.patch.object(render.subprocess, "run", side_effect=convert):
                    try:
                        result = render.to_pdf(src, self.root / "pdf")
                    except (SystemExit, subprocess.TimeoutExpired):
                        self.assertNotEqual(mode, "success")
                    else:
                        self.assertEqual(mode, "success")
                        self.assertEqual(result.read_bytes(), b"synthetic pdf")
                    self.assertEqual(list(scratch.iterdir()), [])

    def test_data_points_stable_and_stops_after_thirty(self):
        text = "\n".join(f"Revenue {i}M" for i in range(1000))
        result = analyze_source.measure(text)
        self.assertEqual(result["data_points"], [f"Revenue {i}M" for i in range(30)])
        original = analyze_source.re.finditer
        count = 0
        def counted(pattern, *args, **kwargs):
            nonlocal count
            for match in original(pattern, *args, **kwargs):
                if pattern.startswith("[^"):
                    count += 1
                yield match
        with mock.patch.object(analyze_source.re, "finditer", side_effect=counted):
            self.assertEqual(analyze_source.measure(text), result)
        self.assertEqual(count, 30)

    def test_compare_selects_numeric_pages(self):
        for directory in [self.root / "source_png", self.root / "out/png"]:
            directory.mkdir(parents=True)
            for page in [9, 10, 99, 100, 999, 1000]:
                Image.new("RGB", (4, 4), (page % 255, 0, 0)).save(directory / f"slide-{page:02d}.png")
        original = Image.open
        opened = []
        def read(path, *args, **kwargs):
            opened.append(Path(path).name)
            return original(path, *args, **kwargs)
        with mock.patch.object(compare_sheet.Image, "open", side_effect=read):
            self.assertEqual(call(compare_sheet, self.root, "--slides", "100,999", "--width", 320)[0], 0)
        self.assertEqual(opened, ["slide-100.png", "slide-100.png", "slide-999.png", "slide-999.png"])


if __name__ == "__main__":
    unittest.main()
