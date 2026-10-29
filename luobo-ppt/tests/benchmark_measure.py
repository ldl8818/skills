"""Print raw scaling observations; deliberately contains no timing assertions."""
import json
from pathlib import Path
import platform
import sys
import time

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from analyze_source import measure
from build_deck import Builder
from common import Style
from layouts import R
from pptx import Presentation

rows = []
for count in (16000, 64000):
    text = "\n".join(f"Revenue {n}M" for n in range(count))
    durations = []
    for _ in range(3):
        start = time.perf_counter()
        result = measure(text)
        durations.append(time.perf_counter() - start)
    rows.append({"records": count, "characters": len(text), "seconds": durations,
                 "result_count": len(result["data_points"]), "data_points": result["data_points"]})
renders = []
for count in (150, 600):
    durations = []
    for _ in range(3):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        builder = Builder(Path("."), Style({}), {}, {}, 13.333, 7.5)
        elements = [R(1, 1, 0.1, 0.1, fill="primary") for _ in range(count)]
        start = time.perf_counter()
        result = builder.render(slide, elements)
        durations.append(time.perf_counter() - start)
        ids = [i for _, group in result for i in group]
        assert len(ids) == len(set(ids)) == count
    renders.append({"shapes": count, "seconds": durations})
print(json.dumps({"python": sys.version, "platform": platform.platform(), "measurements": rows,
                  "native_shape_rendering": renders}, indent=2))
